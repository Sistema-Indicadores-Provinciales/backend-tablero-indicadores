"""Bounded tabular import and explicit chart semantics, independent of HTTP/storage."""
import csv
import math
import re
import unicodedata
from datetime import date, datetime, time
from functools import cmp_to_key
from pathlib import Path
from zipfile import ZipFile, BadZipFile

import numpy as np
import pandas as pd
from openpyxl import load_workbook

MAX_ROWS = 100_000
MAX_COLS = 256
MAX_CELLS = 2_000_000
MAX_POINTS = 10_000
MAX_MAP_POINTS = 2_000
CHART_TYPES = {"bar", "horizontal", "stacked", "line", "area", "scatter", "pie", "donut", "histogram", "box", "heatmap", "table", "indicator", "map"}
MULTI_METRIC_CHART_TYPES = {"bar", "horizontal", "stacked", "line"}
SORT_AGGREGATIONS = {"alphabetical", "sum", "max", "min", "count", "distinct"}

class DataError(ValueError):
    pass

def safe(value):
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    if isinstance(value, (np.integer, int)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        return float(value) if math.isfinite(value) else None
    if isinstance(value, (datetime, date, time, pd.Timestamp)):
        return value.isoformat()
    return str(value)

def bounded(rows):
    result = []
    cells = 0
    for row in rows:
        if len(result) >= MAX_ROWS or len(row) > MAX_COLS:
            raise DataError("La hoja supera el límite de 100.000 filas o 256 columnas. Dividila en hojas más pequeñas.")
        cells += len(row)
        if cells > MAX_CELLS:
            raise DataError("La hoja supera los 2 millones de celdas. Reducí el rango de datos.")
        result.append([safe(v) for v in row])
    return result

def inspect_zip(path):
    try:
        with ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:
                raise DataError("El Excel descomprimido supera 100 MB.")
    except BadZipFile as exc:
        raise DataError("El archivo no es un Excel válido o está protegido con contraseña.") from exc

def read_rows(path, sheet=None):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        inspect_zip(path)
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            names = book.sheetnames
            if sheet is None:
                return names
            if sheet not in names:
                raise DataError("La hoja seleccionada no existe.")
            ws = book[sheet]
            if ws.max_row > MAX_ROWS or ws.max_column > MAX_COLS:
                raise DataError("La hoja supera los límites de filas o columnas.")
            return bounded(ws.iter_rows(values_only=True))
        finally:
            book.close()
    if suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            if sheet is None:
                return book.sheet_names()
            ws = book.sheet_by_name(sheet)
            if ws.nrows > MAX_ROWS or ws.ncols > MAX_COLS or ws.nrows * ws.ncols > MAX_CELLS:
                raise DataError("La hoja supera los límites de lectura.")
            def rows():
                for i in range(ws.nrows):
                    yield [xlrd.xldate_as_datetime(c.value, book.datemode) if c.ctype == xlrd.XL_CELL_DATE else None if c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_ERROR) else c.value for c in ws.row(i)]
            return bounded(rows())
        finally:
            book.release_resources()
    if suffix in {".csv", ".tsv"}:
        if sheet is None:
            return ["Datos"]
        if sheet != "Datos":
            raise DataError("La hoja seleccionada no existe.")
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                with path.open(encoding=encoding, newline="") as stream:
                    sample = stream.read(8192)
                    stream.seek(0)
                    try:
                        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
                    except csv.Error:
                        dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
                    return bounded(csv.reader(stream, dialect))
            except UnicodeDecodeError:
                continue
        raise DataError("No se pudo leer la codificación del CSV. Guardalo como UTF-8.")
    raise DataError("Formato admitido: XLSX, XLSM, XLS, CSV o TSV.")

def number(value, decimal=","):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return safe(float(value))
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    percent = text.endswith("%")
    text = text.rstrip("%")
    text = re.sub(r"^(?:ARS|USD|EUR|[$€£])", "", text, flags=re.I)
    thousands = "." if decimal == "," else ","
    # A thousands separator must group exactly three digits. Do not turn 1.5 into 15.
    integer = text.split(decimal)[0].lstrip("+-")
    if thousands in text:
        if not re.fullmatch(r"\d{1,3}(?:" + re.escape(thousands) + r"\d{3})+", integer):
            return None
        if decimal in text and thousands in text.split(decimal, 1)[1]:
            return None
    text = text.replace(thousands, "").replace(decimal, ".")
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", text):
        return None
    result = float(text) * (-1 if negative else 1) / (100 if percent else 1)
    return result if math.isfinite(result) else None

def table(rows, header=1, decimal=",", types=None):
    if not rows:
        raise DataError("La hoja está vacía.")
    if header < 0 or header > len(rows):
        raise DataError("La fila de encabezados no existe. Usá 0 si no hay encabezados.")
    width = max(map(len, rows))
    raw_names = rows[header - 1] if header else []
    names, used = [], set()
    for i in range(width):
        name = str(raw_names[i]).strip() if i < len(raw_names) and raw_names[i] is not None else ""
        base = name or f"Columna {i + 1}"
        name, n = base, 2
        while name in used:
            name, n = f"{base} ({n})", n + 1
        names.append(name)
        used.add(name)
    records = [list(r) + [None] * (width - len(r)) for r in rows[header:]]
    records = [r for r in records if any(v is not None and v != "" for v in r)]
    if not records:
        raise DataError("No hay datos debajo de los encabezados. Revisá la fila elegida.")
    frame = pd.DataFrame(records, columns=names, dtype=object)
    warnings = []
    meta = {}
    types = types or {}
    if set(types) - set(names):
        raise DataError("La configuración de tipos contiene columnas inexistentes.")
    for col in names:
        values = frame[col].tolist()
        present = [v for v in values if v is not None and v != ""]
        kind = types.get(col, "auto")
        if kind == "auto":
            numeric = present and all(number(v, decimal) is not None for v in present)
            ids = any(isinstance(v, str) and re.match(r"^0\d+$", v) for v in present)
            kind = "boolean" if present and all(isinstance(v, bool) for v in present) else "number" if numeric and not ids else "text"
            if kind == "text" and present and all(re.match(r"^\d{4}-\d{2}-\d{2}(?:T.*)?$", str(v)) for v in present):
                kind = "date"
        if kind == "number":
            converted = [number(v, decimal) for v in values]
        elif kind == "date":
            def parse_date(v):
                if v is None or v == "":
                    return None
                try:
                    text = str(v)
                    if re.match(r"^\d{4}-\d{2}-\d{2}", text):
                        return pd.to_datetime(text, format="ISO8601").isoformat()
                    return pd.to_datetime(text, dayfirst=decimal == ",").isoformat()
                except (ValueError, TypeError, OverflowError):
                    return None
            converted = [parse_date(v) for v in values]
        elif kind == "boolean":
            mapping = {"true": True, "false": False, "si": True, "sí": True, "no": False, "1": True, "0": False}
            converted = [mapping.get(str(v).strip().lower()) if v is not None else None for v in values]
        elif kind == "text":
            converted = [str(v) if v is not None else None for v in values]
        else:
            raise DataError("Tipo de columna no válido.")
        invalid = sum(v is not None and v != "" and c is None for v, c in zip(values, converted))
        if invalid:
            warnings.append(f"{col}: {invalid} valores no pudieron convertirse a {kind}; se tratarán como vacíos.")
        frame[col] = pd.Series(converted, dtype=object)
        unique = list(dict.fromkeys(v for v in converted if v is not None))
        meta[col] = {"type": kind, "unique_values": unique[:100], "unique_count": len(unique), "missing": sum(v is None for v in converted), "invalid": invalid}
    return frame, {"columns": names, "column_meta": meta, "row_count": len(frame), "preview": [{c: safe(v) for c, v in row.items()} for row in frame.head(20).to_dict("records")], "warnings": warnings}

def filter_value(value):
    value = safe(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def filter_keys(series, values):
    sample = next((safe(v) for v in series if safe(v) is not None), None)
    selected = [str(v) for v in values]
    if isinstance(sample, (int, float)) and not isinstance(sample, bool):
        def numeric_key(value):
            try:
                parsed = float(value)
                return filter_value(parsed) if math.isfinite(parsed) else value
            except (ValueError, TypeError, OverflowError):
                return value
        selected = [numeric_key(v) for v in selected]
    elif isinstance(sample, bool):
        selected = [{'true': 'True', 'false': 'False'}.get(v.lower(), v) for v in selected]
    return series.map(filter_value), selected


def filter_selection(selection):
    """Normalize a legacy inclusion list or a bounded include/exclude selection."""
    if isinstance(selection, list):
        return "include", selection
    if isinstance(selection, dict) and len(selection) == 1:
        mode, values = next(iter(selection.items()))
        if mode in {"include", "exclude"} and isinstance(values, list):
            return mode, values
    raise DataError("Filtro no válido.")


def filter_frame(frame, filters):
    for col, values in filters.items():
        if col not in frame or not isinstance(values, list):
            raise DataError("Filtro no válido.")
        if values:
            keys, selected = filter_keys(frame[col], values)
            frame = frame[keys.isin(selected)]
    return frame


def natural_key(value):
    """Build a stable, accent-insensitive key that keeps numeric values natural."""
    value = safe(value)
    if value is None:
        return (1, ())
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return (0, ((0, float(value)),))
    normalized = unicodedata.normalize("NFD", str(value).casefold())
    normalized = "".join(char for char in normalized if unicodedata.category(char) != "Mn")
    chunks = tuple((0, int(part)) if part.isdigit() else (1, part) for part in re.split(r"(\d+)", normalized) if part)
    return (0, chunks)


def _sort_score(values, aggregation):
    present = [safe(value) for value in values if safe(value) is not None]
    if aggregation == "alphabetical":
        return min(present, key=natural_key) if present else None
    if aggregation == "count":
        return len(present)
    if aggregation == "distinct":
        return len({json_key(value) for value in present})
    if any(not isinstance(value, (int, float, bool)) for value in present):
        raise DataError("El campo elegido para ordenar debe estar configurado como número.")
    numeric = [float(value) for value in present]
    if not numeric:
        return None
    if aggregation == "sum":
        return sum(numeric)
    if aggregation == "max":
        return max(numeric)
    if aggregation == "min":
        return min(numeric)
    raise DataError("Operación de orden no válida.")


def json_key(value):
    value = safe(value)
    return f"{type(value).__name__}:{value}"


def _compare_sort_items(first, second, direction):
    first_value, first_score = first
    second_value, second_score = second
    if first_score is None and second_score is not None:
        return 1
    if second_score is None and first_score is not None:
        return -1
    first_key, second_key = natural_key(first_score), natural_key(second_score)
    comparison = (first_key > second_key) - (first_key < second_key)
    if direction == "desc":
        comparison *= -1
    if comparison:
        return comparison
    first_name, second_name = natural_key(first_value), natural_key(second_value)
    return (first_name > second_name) - (first_name < second_name)


def sort_group_values(frame, group_column, config, *, strict=True):
    """Order dimension values using a selected field and aggregation."""
    if group_column not in frame:
        raise DataError("El campo que se quiere ordenar no existe.")
    groups = frame.groupby(group_column, dropna=False, sort=False)
    values = [safe(value) for value, _ in groups]
    if not config:
        return sorted(values, key=natural_key)
    aggregation = config.get("aggregation", "alphabetical")
    direction = config.get("direction", "asc")
    field = config.get("field") or group_column
    if aggregation not in SORT_AGGREGATIONS or direction not in {"asc", "desc"}:
        raise DataError("Configuración de orden no válida.")
    if field not in frame:
        if strict:
            raise DataError(f"El campo elegido para ordenar, '{field}', no existe en estos datos.")
        return values
    scored = [(safe(value), _sort_score(subframe[field].tolist(), aggregation)) for value, subframe in groups]
    return [value for value, _ in sorted(scored, key=cmp_to_key(lambda a, b: _compare_sort_items(a, b, direction)))]


def contextual_filter_options(frame, columns, filters, sorts=None):
    # Normalize once per column. Counting failed conditions lets each facet exclude
    # its own selection without refiltering/downloading the table for every field.
    keyed = {col: filter_keys(frame[col], filter_selection(filters.get(col, []))[1]) for col in frame.columns if col in columns}
    misses = {}
    for col, (keys, selected) in keyed.items():
        mode, values = filter_selection(filters.get(col, []))
        if not values:
            continue
        matches = keys.isin(selected)
        misses[col] = (~matches if mode == "include" else matches).astype('int16')
    failures = sum(misses.values()) if misses else None
    options = {}
    sorts = sorts or {}
    for col, (keys, selected) in keyed.items():
        mask = None if failures is None else (failures - misses.get(col, 0)) == 0
        visible = keys if mask is None else keys[mask]
        visible_frame = frame if mask is None else frame[mask]
        sort_config = sorts.get(col)
        sort_applied = bool(sort_config and (sort_config.get("field") or col) in visible_frame.columns)
        if sort_applied:
            ordered = sort_group_values(visible_frame, col, sort_config, strict=False)
            values = [filter_value(value) for value in ordered]
        else:
            values = list(dict.fromkeys(visible))
        available = set(values)
        _, raw_selected = filter_selection(filters.get(col, []))
        options[col] = {
            'values': values[:100], 'total': len(values), 'unfiltered_total': keys.nunique(),
            'sort_applied': sort_applied,
            # Validate selected values against the whole facet, beyond the 100 suggestions.
            'available_selected': [value for value, key in zip(raw_selected, selected) if key in available],
        }
    return options


def apply_view_filters(frame, filters):
    """Apply section filters, preserving list-based inclusion requests."""
    result = frame
    for col, selection in filters.items():
        if col not in result:
            continue
        mode, values = filter_selection(selection)
        if not values:
            continue
        keys, selected = filter_keys(result[col], values)
        matches = keys.isin(selected)
        result = result[~matches if mode == "exclude" else matches]
    return result


def bucket_dates(frame, column, bucket):
    """Replace a date dimension with the requested calendar bucket."""
    if bucket == "none":
        return
    if bucket not in {"month", "year"}:
        raise DataError("Agrupación temporal no válida.")
    dates = pd.to_datetime(frame[column], errors="coerce")
    if dates.isna().any():
        raise DataError("Para agrupar por fecha, configurá el eje X como fecha y corregí los valores vacíos.")
    frame[column] = dates.dt.strftime("%Y-%m" if bucket == "month" else "%Y")


def coordinate(value, limit, decimal=","):
    """Parse decimal degrees, including coordinates split into two 3-digit blocks."""
    if value is None or isinstance(value, bool):
        return None, False
    parsed = number(value, decimal)
    normalized = False
    if parsed is None:
        text = str(value).strip().replace("\u00a0", "").replace(" ", "")
        match = re.fullmatch(r"([+-]?)(\d{1,3})([,.])(\d{3})\3(\d{3})", text)
        if not match:
            return None, False
        parsed = float(f"{match.group(1)}{match.group(2)}.{match.group(4)}{match.group(5)}")
        normalized = True
    parsed = float(parsed)
    if abs(parsed) > limit and abs(parsed) >= 1_000_000 and parsed.is_integer():
        scaled = parsed / 1_000_000
        if abs(scaled) <= limit:
            parsed = scaled
            normalized = True
    if not math.isfinite(parsed) or abs(parsed) > limit:
        return None, normalized
    return parsed, normalized


def chart(frame, cfg):
    kind = cfg.get("chart_type", "bar")
    agg = cfg.get("aggregation", "sum")
    x, y, group = cfg.get("x_col", ""), cfg.get("y_col", ""), cfg.get("group_col", "")
    additional_metrics = cfg.get("y_cols") or []
    metrics = list(dict.fromkeys(column for column in [y, *additional_metrics] if column))
    multiple_metrics = kind in MULTI_METRIC_CHART_TYPES and agg != "count" and len(metrics) > 1
    active_metrics = metrics if multiple_metrics else [y]
    if kind not in CHART_TYPES or agg not in {"sum", "mean", "median", "min", "max", "count", "distinct"}:
        raise DataError("Tipo de gráfico u operación no válido.")
    if kind == "map":
        latitude = cfg.get("latitude_col", "")
        longitude = cfg.get("longitude_col", "")
        label = cfg.get("label_col", "")
        size = cfg.get("size_col", "")
        map_metrics = list(dict.fromkeys(cfg.get("map_metric_cols") or []))
        mode = cfg.get("map_mode", "markers")
        required = [latitude, longitude] + ([label] if label else []) + ([size] if size else []) + map_metrics
        if mode not in {"markers", "circles"} or any(column not in frame for column in required):
            raise DataError("Elegí columnas válidas de latitud y longitud para el mapa.")
        frame = filter_frame(frame.copy(), cfg.get("filters", {}))
        rows = len(frame)
        if rows > MAX_MAP_POINTS:
            raise DataError("El mapa supera 2.000 puntos. Aplicá filtros antes de generarlo.")
        points, invalid_coordinates, invalid_sizes, normalized_coordinates = [], 0, 0, 0
        for index, row in frame.iterrows():
            lat, normalized_latitude = coordinate(row[latitude], 90, cfg.get("decimal", ","))
            lng, normalized_longitude = coordinate(row[longitude], 180, cfg.get("decimal", ","))
            if lat is None or lng is None:
                invalid_coordinates += 1
                continue
            normalized_coordinates += int(normalized_latitude) + int(normalized_longitude)
            value = safe(row[size]) if size else None
            if value is not None and (not isinstance(value, (int, float)) or isinstance(value, bool)):
                invalid_sizes += 1
                value = None
            points.append({
                "id": str(index), "lat": float(lat), "lng": float(lng),
                "label": str(safe(row[label])) if label and safe(row[label]) is not None else f"Punto {len(points) + 1}",
                "value": value,
                **({"metrics": [{"field": metric, "value": safe(row[metric])} for metric in map_metrics]} if map_metrics else {}),
            })
        if rows and not points:
            raise DataError("No hay coordenadas válidas para mostrar en el mapa.")
        warnings = []
        if invalid_coordinates:
            warnings.append(f"Se omitieron {invalid_coordinates} filas con coordenadas inválidas.")
        if invalid_sizes:
            warnings.append(f"{invalid_sizes} valores de tamaño no eran numéricos; se usará el tamaño predeterminado.")
        if normalized_coordinates:
            warnings.append(f"Se normalizaron {normalized_coordinates} coordenadas escritas con seis decimales separados en bloques.")
        return {"labels": [], "datasets": [], "map_points": points, "filtered_rows": rows, "warnings": warnings}
    if kind == "table" and additional_metrics:
        required_metrics = metrics
    else:
        required_metrics = active_metrics if agg != "count" or kind in {"scatter", "box", "histogram"} else []
    required = [x, *required_metrics] + ([group] if group else [])
    if any(c not in frame for c in required):
        raise DataError("Elegí columnas válidas para este gráfico.")
    frame = frame.copy()
    frame = filter_frame(frame, cfg.get("filters", {}))
    rows = len(frame)
    warnings = []
    if kind == "table":
        # Las tablas antiguas mantienen todas las columnas; al agregar métricas,
        # la selección pasa a definir explícitamente sus columnas visibles.
        cols = list(dict.fromkeys([x, *metrics] + ([group] if group else []))) if additional_metrics else list(frame.columns)
        return {"labels": [], "datasets": [], "filtered_rows": rows, "columns": cols, "records": [{c: safe(v) for c, v in r.items()} for r in frame[cols].head(500).to_dict("records")], "warnings": ["Vista limitada a 500 filas."] if rows > 500 else []}
    if not rows:
        return {"labels": [], "datasets": [], "filtered_rows": 0, "warnings": []}
    if multiple_metrics:
        numeric_required = agg not in {"distinct"}
        if numeric_required:
            for metric in active_metrics:
                if any(value is not None and not isinstance(value, (int, float, bool)) for value in frame[metric]):
                    raise DataError(f"'{metric}' contiene texto. Elegí contar valores distintos o configurá su tipo como número en la vista previa.")
                missing = frame[metric].isna().sum()
                if missing:
                    warnings.append(f"{missing} filas sin valor numérico no se usaron para {metric}.")
            # Una fila puede contribuir a una métrica aunque otra esté vacía.
            aggregate_frame = frame[frame[active_metrics].notna().any(axis=1)].copy()
            if aggregate_frame.empty:
                raise DataError("No hay valores numéricos válidos para esta selección.")
        else:
            aggregate_frame = frame.copy()
        bucket_dates(aggregate_frame, x, cfg.get("date_bucket", "none"))
        keys = list(dict.fromkeys([x] + ([group] if group else [])))
        grouped = aggregate_frame.groupby(keys, dropna=False, sort=False)
        results = {}
        for metric in active_metrics:
            if agg == "distinct":
                results[metric] = grouped[metric].nunique()
            elif agg == "sum":
                results[metric] = grouped[metric].sum(min_count=1)
            else:
                results[metric] = getattr(grouped[metric], agg)()
        labels = sort_group_values(aggregate_frame, x, cfg.get("category_sort"))
        series_label = lambda value: "Sin dato" if safe(value) is None else str(safe(value))
        if group:
            group_values = [series_label(value) for value in sort_group_values(aggregate_frame, group, cfg.get("series_sort"))]
        else:
            group_values = [None]
        series_count = len(active_metrics) * len(group_values)
        if series_count > 50:
            raise DataError("Hay más de 50 series entre métricas y grupos. Aplicá filtros o quitá métricas.")
        if len(labels) * series_count > MAX_POINTS:
            raise DataError("El resultado supera 10.000 puntos. Aplicá filtros, quitá métricas o agrupá por mes/año.")
        group_index = keys.index(group) if group else None
        datasets = []
        for metric in active_metrics:
            pairs = [(key if isinstance(key, tuple) else (key,), safe(value)) for key, value in results[metric].items()]
            matrix = {(safe(key[0]), series_label(key[group_index]) if group else None): value for key, value in pairs}
            for group_value in group_values:
                datasets.append({
                    "label": f"{metric} · {group_value}" if group else metric,
                    "metric": metric,
                    **({"group": group_value} if group else {}),
                    "data": [matrix.get((label, group_value)) for label in labels],
                })
        return {"labels": labels, "datasets": datasets, "filtered_rows": rows, "warnings": warnings}
    numeric_required = agg not in {"count", "distinct"} or kind in {"scatter", "histogram", "box"}
    if numeric_required:
        if any(v is not None and not isinstance(v, (int, float, bool)) for v in frame[y]):
            raise DataError(f"'{y}' contiene texto. Elegí contar registros o configurá su tipo como número en la vista previa.")
        missing = frame[y].isna().sum()
        if missing:
            warnings.append(f"Se omitieron {missing} filas sin valor numérico en {y}.")
        frame = frame[frame[y].notna()]
        if frame.empty:
            raise DataError("No hay valores numéricos válidos para esta selección.")
    bucket_dates(frame, x, cfg.get("date_bucket", "none"))
    if kind in {"scatter", "histogram", "box"}:
        if len(frame) > MAX_POINTS:
            raise DataError("El gráfico supera 10.000 puntos. Aplicá filtros antes de generar.")
        if kind == "scatter" and any(v is None or not isinstance(v, (int, float)) for v in frame[x]):
            raise DataError("La dispersión necesita un eje X numérico sin vacíos.")
        datasets = []
        grouping = group or (x if kind == "box" else None)
        groups = frame.groupby(grouping, dropna=False, sort=False) if grouping else [(y, frame)]
        for label, sub in groups:
            label = safe(label)
            datasets.append({"label": "Sin dato" if label is None else str(label), "x": [safe(v) for v in sub[x]], "data": [safe(v) for v in sub[y]]})
        if grouping and cfg.get("series_sort"):
            ordered = [str(value if value is not None else "Sin dato") for value in sort_group_values(frame, grouping, cfg["series_sort"])]
            ranks = {value: index for index, value in enumerate(ordered)}
            datasets.sort(key=lambda dataset: ranks.get(dataset["label"], len(ranks)))
        if len(datasets) > 50:
            raise DataError("Hay más de 50 series. Aplicá filtros.")
        return {"labels": [], "datasets": datasets, "filtered_rows": rows, "warnings": warnings}
    keys = list(dict.fromkeys([x] + ([group] if group else [])))
    grouped = frame.groupby(keys, dropna=False, sort=False)
    if agg == "count":
        result = grouped.size()
    elif agg == "distinct":
        result = grouped[y].nunique()
    elif agg == "sum":
        result = grouped[y].sum(min_count=1)
    else:
        result = getattr(grouped[y], agg)()
    if len(result) > MAX_POINTS:
        raise DataError("El resultado supera 10.000 puntos. Aplicá filtros o agrupá por mes/año.")
    pairs = [(k if isinstance(k, tuple) else (k,), safe(v)) for k, v in result.items()]
    labels = sort_group_values(frame, x, cfg.get("category_sort"))
    group_index = keys.index(group) if group else None
    series_label = lambda value: "Sin dato" if safe(value) is None else str(safe(value))
    series = list(dict.fromkeys(series_label(k[group_index]) for k, _ in pairs)) if group else ["Registros" if agg == "count" else y]
    if group and cfg.get("series_sort"):
        series = [series_label(value) for value in sort_group_values(frame, group, cfg["series_sort"])]
    if len(series) > 50:
        raise DataError("Hay más de 50 series. Aplicá filtros.")
    matrix = {(safe(k[0]), series_label(k[group_index]) if group else series[0]): v for k, v in pairs}
    datasets = [{"label": s, "data": [matrix.get((v, s)) for v in labels]} for s in series]
    if kind in {"pie", "donut"}:
        if len(datasets) > 1 or len(labels) > 30:
            raise DataError("Circular/anillo admite una serie y hasta 30 categorías. Quitá la agrupación o aplicá filtros.")
        values = datasets[0]["data"]
        if any(v is not None and v < 0 for v in values) or not any(v and v > 0 for v in values):
            raise DataError("Circular/anillo necesita valores no negativos y al menos uno positivo.")
    if kind == "indicator":
        if agg == "count":
            value = rows
        elif agg == "distinct":
            value = int(frame[y].nunique())
        elif agg == "sum":
            value = safe(frame[y].sum(min_count=1))
        else:
            value = safe(getattr(frame[y], agg)())
        return {"labels": ["Total"], "datasets": [{"label": agg, "data": [value]}], "filtered_rows": rows, "warnings": warnings}
    return {"labels": labels, "datasets": datasets, "filtered_rows": rows, "warnings": warnings}
